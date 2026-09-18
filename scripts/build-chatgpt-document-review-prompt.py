#!/usr/bin/env python3
"""Build a self-contained independent-review prompt from the generated DOCX pack."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
PACK = ROOT / "output" / "application-documents-2026-09-12"
OUTPUT = ROOT / "tmp" / "chatgpt-document-review-prompt.txt"

FILES = [
    ("Current CV", "Christopher_Ongko_CV_2026-09.docx", 2),
    ("ERA:AI Winter 2027 response bank", "ERA_AI_Winter_2027_Response_Bank.docx", 2),
    ("HKU AI Engineer / RA II packet", "HKU_AI_Engineer_RAII_Application_Packet.docx", 2),
    ("VU Social Data Science letter", "VU_Social_Data_Science_Application_Letter.docx", 1),
    ("VU ATLANTIS computational-track letter", "VU_ATLANTIS_Computational_Track_Cover_Letter.docx", 1),
    ("NHH statement of purpose", "NHH_Finance_Statement_of_Purpose.docx", 1),
    ("NHH tentative proposal", "NHH_Finance_Tentative_Research_Proposal.docx", 5),
    ("NHH publications/research-activities list", "NHH_Publications_and_Research_Activities.docx", 1),
]


def extract(path: Path) -> str:
    doc = Document(path)
    blocks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in doc.tables:
        blocks.append("[TABLE]")
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
    return "\n".join(blocks)


def main() -> None:
    parts = [
        "You are the independent QA reviewer for an evidence-bounded application-document pack. Do not merely praise the writing. Audit whether each route is actually ready for browser drafting or still blocked.",
        "",
        "Current date: 2026-09-12. Applicant: Christopher Ongko, currently completing a finance master’s at Yuan Ze University. No application has been submitted.",
        "",
        "Official sources to verify:",
        "- ERA:AI Winter 2027: https://erafellowship.org/fellowship",
        "- HKU Ref. 537095: https://jobs.hku.hk/cw/en/job/537095/",
        "- VU Social Data Science: https://workingat.vu.nl/vacancies/phd-position-in-social-data-science-amsterdam-1330114",
        "- VU ATLANTIS: https://workingat.vu.nl/vacancies/three-fully-funded-phd-positions-in-competition-law-and-ai-amsterdam-1307827",
        "- NHH PhD admission: https://www.nhh.no/en/study-programmes/phd-programme-at-nhh/admission/",
        "",
        "Review instructions:",
        "1. Cross-check factual consistency across all eight documents: identity, degree state, dates, paper titles/status, project claims, evidence boundaries, and corrected $170B Invisible Ledger framing.",
        "2. Cross-check each packet against the live official requirements and document/page/word limits. Separate application copy from mandatory formal attachments.",
        "3. Flag any statement that is unsupported, strategically self-defeating, too generic, too AI-polished, or likely to trigger an avoidable rejection. Do not recommend hiding genuine eligibility problems.",
        "4. For each route, return exactly one readiness state: READY_FOR_BROWSER_DRAFT, READY_AFTER_SMALL_DOCUMENT_EDIT, WAITING_FORMAL_RECORDS, ELIGIBILITY_BLOCKED_OR_UNCERTAIN, or NOT_ENOUGH_INFORMATION.",
        "5. Give the exact small edits required before use. Distinguish draftable document gaps from non-generatable blockers such as transcripts, degree completion, referees, teammate commitments, fees, work authorization, attestations, CAPTCHA/OTP, and final submit.",
        "6. Check packaging quality using these verified final render counts: CV 2 pages; ERA 2; HKU 2; VU Social 1; VU ATLANTIS 1; NHH SOP 1 (252 words); NHH proposal 5 (2,212 words); NHH research list 1. All pages were rendered in LibreOffice and visually inspected locally.",
        "7. End with a compact matrix: route, document completeness, eligibility state, remaining hard gates, and verdict. Be candid and evidence-sensitive.",
        "",
        "DOCUMENTS FOLLOW",
    ]
    for label, filename, pages in FILES:
        path = PACK / filename
        parts.extend([
            "",
            f"===== {label} | {filename} | verified render: {pages} page(s) =====",
            extract(path),
        ])
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text("\n".join(parts) + "\n", encoding="utf-8")
    print(OUTPUT)


if __name__ == "__main__":
    main()
